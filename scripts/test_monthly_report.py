from datetime import date
import os
import unittest

os.environ.setdefault(
    'DATABASE_URL',
    'postgresql://test:test@example.invalid:5432/postgres',
)
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'control_asistencia.settings')

import django
django.setup()

import openpyxl
from io import BytesIO
from unittest.mock import patch
from django.urls import resolve, reverse
from app.services import ReporteService
from app.views import _nombre_hoja_empleado, exportar_reporte_mensual_empleados


class MonthlyReportTests(unittest.TestCase):
    def test_weeks_are_all_monday_to_friday_blocks(self):
        ranges = ReporteService.rangos_semanales_mes(date(2026, 8, 31))
        self.assertEqual(
            ranges,
            [
                (1, date(2026, 8, 3), date(2026, 8, 7)),
                (2, date(2026, 8, 10), date(2026, 8, 14)),
                (3, date(2026, 8, 17), date(2026, 8, 21)),
                (4, date(2026, 8, 24), date(2026, 8, 28)),
                (5, date(2026, 8, 31), date(2026, 9, 4)),
            ],
        )

    def test_february_uses_all_work_weeks(self):
        ranges = ReporteService.rangos_semanales_mes(date(2024, 2, 10))
        self.assertEqual(ranges[-1], (4, date(2024, 2, 26), date(2024, 3, 1)))

    def test_excel_sheet_names_are_valid_and_unique(self):
        empleado = type('Empleado', (), {
            'nombre_completo': 'Nombre / muy largo : para [Excel] empleado repetido',
            'pk': 1,
        })()
        usados = set()
        primero = _nombre_hoja_empleado(empleado, usados)
        segundo = _nombre_hoja_empleado(empleado, usados)
        self.assertLessEqual(len(primero), 31)
        self.assertNotEqual(primero, segundo)
        self.assertFalse(any(char in primero for char in '[]:*?/\\'))

    @patch('app.views.RegistroAsistencia.objects')
    @patch('app.views.Empleado.objects')
    def test_empty_database_still_generates_valid_workbook(self, empleados, registros):
        empleados.order_by.return_value = []
        registros.filter.return_value.select_related.return_value.order_by.return_value = []
        undecorated = exportar_reporte_mensual_empleados.__wrapped__
        response = undecorated(object())
        workbook = openpyxl.load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Sin empleados'])
        self.assertIn('No hay empleados', workbook.active['A1'].value)

    @patch('app.views.timezone.localdate', return_value=date(2026, 8, 15))
    @patch('app.views.RegistroAsistencia.objects')
    @patch('app.views.Empleado.objects')
    def test_august_layout_matches_reference(self, empleados, registros, _localdate):
        empleado = type('Empleado', (), {
            'nombre_completo': 'Ana Prueba', 'id_empleado': 1, 'pk': 1,
        })()
        empleados.order_by.return_value = [empleado]
        registros.filter.return_value.select_related.return_value.order_by.return_value = []
        request = type('Request', (), {'GET': {'mes': '2026-08'}})()
        response = exportar_reporte_mensual_empleados.__wrapped__(request)
        sheet = openpyxl.load_workbook(BytesIO(response.content)).active
        week_rows = [
            row[0] for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True)
            if row[0] and str(row[0]).startswith('Semana ')
        ]
        self.assertEqual(week_rows, [
            'Semana 1 (03/08/2026 - 07/08/2026)',
            'Semana 2 (10/08/2026 - 14/08/2026)',
            'Semana 3 (17/08/2026 - 21/08/2026)',
            'Semana 4 (24/08/2026 - 28/08/2026)',
            'Semana 5 (31/08/2026 - 04/09/2026)',
        ])
        dates = [
            row[0] for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True)
            if isinstance(row[0], str) and row[0][:2].isdigit()
        ]
        self.assertEqual(len(dates), 25)

    def test_existing_attendance_download_uses_monthly_report(self):
        url = reverse('descargar_excel')
        self.assertEqual(url.rstrip('/'), '/login/descargar/asistencia')
        self.assertIs(resolve(url).func, exportar_reporte_mensual_empleados)

    @patch('app.views.RegistroAsistencia.objects')
    @patch('app.views.Empleado.objects')
    def test_requested_august_is_used_even_when_current_month_changes(self, empleados, registros):
        empleados.order_by.return_value = []
        registros.filter.return_value.select_related.return_value.order_by.return_value = []
        request = type('Request', (), {'GET': {'mes': '2026-08'}})()
        response = exportar_reporte_mensual_empleados.__wrapped__(request)
        self.assertIn('2026_08', response['Content-Disposition'])
        registros.filter.assert_called_once_with(
            fecha_registro__range=(date(2026, 8, 1), date(2026, 8, 31))
        )


if __name__ == '__main__':
    unittest.main()
